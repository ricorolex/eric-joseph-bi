"""Board-ready memo export (.docx and .pdf) built from the dashboard's own numbers.

Sections: executive summary (the AI brief, when one has been generated), macro scorecard, sector breakdown,
banking micro-analysis (small businesses vs FDIC branches, deposit capture, concentration), compliance note.
"""
from __future__ import annotations

import io
import re
from datetime import datetime

import numpy as np
import pandas as pd

from banking import branch_outliers, bank_share, catchment_branches

NAVY, TEAL = "0B1F3A", "0E7490"
DISCLAIMER = (
    "Screening support only; not legal, financial or investment advice. Figures come from public data (U.S. Census "
    "Bureau ACS and Business Patterns, Bureau of Labor Statistics QCEW, Bureau of Economic Analysis, FDIC Summary of "
    "Deposits) and should be verified before acting. Small-business counts are employer establishments (locations), "
    "not companies, and exclude sole proprietors. Branch deposits are as booked at each branch and can be distorted "
    "by headquarters or online branches. This memo contains no race, ethnicity or other protected-class data and "
    "must not be used to include or exclude communities from lending, housing or other services on those bases "
    "(fair-lending and fair-housing laws apply)."
)


# ------------------------------------------------------------------ assemble inputs
def _fmt_money(v) -> str:
    return "n/a" if pd.isna(v) else f"${v:,.0f}"


def _fmt_num(v, d=0, suffix="") -> str:
    return "n/a" if pd.isna(v) else f"{v:,.{d}f}{suffix}"


def collect(view: pd.DataFrame, sector_lq: pd.DataFrame, selected_sector: str, weights_pct: dict, filters: dict,
            loaded_at: str, status: pd.DataFrame, banking_ctx: dict | None, brief_text: str | None,
            purpose: str, prepared_for: str = "", top_n: int = 8) -> dict:
    v = view.head(top_n)
    score = pd.DataFrame({
        "Rank": v["rank"].astype(int).astype(str),
        "Market": v["city"], "ZIP": v["zip_code"],
        "Score": v["BI_Score"].map(lambda x: f"{x:.0f}"),
        "Sector LQ": v["lq"].map(lambda x: _fmt_num(x, 2)),
        "Median income": v["median_income"].map(_fmt_money),
        "GDP growth": v["gdp_growth_pct"].map(lambda x: _fmt_num(x, 1, "%")),
        "Bachelor's+": v["highly_educated_pct"].map(lambda x: _fmt_num(x, 0, "%")),
    })
    top_zip = str(view.iloc[0]["zip_code"])
    sec = sector_lq[sector_lq["zip_code"] == top_zip].sort_values("lq", ascending=False).head(6)
    sector_table = pd.DataFrame({"Sector": sec["sector"], "Location quotient": sec["lq"].map(lambda x: _fmt_num(x, 2))})
    m = {
        "title": "Market Investment Committee Memo",
        "prepared_for": prepared_for.strip(), "purpose": purpose,
        "date": datetime.now().strftime("%B %d, %Y").replace(" 0", " "),
        "data_as_of": loaded_at, "summary_md": (brief_text or "").strip() or None,
        "settings": [f"Selected sector: {selected_sector}",
                     "Score weights: " + ", ".join(f"{k.replace('_', ' ')} {v_}%" for k, v_ in weights_pct.items()),
                     f"Filters: minimum median income ${filters['min_median_income']:,}; minimum bachelor's "
                     f"{filters['min_bachelors_only_pct']}%"],
        "scorecard": score, "sector_title": f"Sector mix in {view.iloc[0]['city']} (highest location quotients)",
        "sector_table": sector_table, "banking": None, "competitors": None, "outliers": None,
        "sources": [f"{r.Source}: {r.State}" for r in status[["Source", "State"]].itertuples()],
        "sample": bool((status["State"] != "Live").any()),
    }
    if banking_ctx:
        t = banking_ctx["table"].head(top_n)
        m["banking"] = pd.DataFrame({
            "ZIP": t["zip_code"], "Market": t["city"], "Small biz": t["small_biz"].map(lambda x: _fmt_num(x)),
            "Branches": t["branches_in_zip"].map(lambda x: _fmt_num(x)),
            "Biz / branch": t["small_biz_per_branch"].map(lambda x: _fmt_num(x)),
            "Coverage": t["coverage"], "Capture idx": t["capture_index"].map(lambda x: _fmt_num(x)),
            "Concentration": t["hhi_label"],
        })
        m["banking_note"] = (f"Small business = {banking_ctx['small_def'].lower()}. Competitor area: "
                             f"{banking_ctx.get('catch_label', str(banking_ctx['radius']) + ' mi')}. "
                             "Capture index = deposits per small business vs the median ZIP (100 = median).")
        top = banking_ctx["table"].iloc[0]
        ring = (banking_ctx.get("catchments") or {}).get(top["zip_code"])
        nb = catchment_branches(banking_ctx["branches"], top["latitude"], top["longitude"], banking_ctx["radius"], ring)
        if not nb.empty:
            sh = bank_share(nb).head(6)
            m["competitors_title"] = f"Deposit share by bank around {top['city']} ({top['zip_code']})"
            m["competitors"] = pd.DataFrame({"Bank": sh["bank"], "Branches": sh["branches"].astype(str),
                                             "Deposits": (sh["deposits_usd"] / 1e6).map(lambda x: f"${x:,.1f}M"),
                                             "Share": sh["share_pct"].map(lambda x: f"{x:.1f}%")})
            o = branch_outliers(nb)
            o = o[o["flag"] == "Underperforming"].sort_values("vs_peer").head(5)
            if not o.empty:
                m["outliers_title"] = f"Underperforming branches near {top['city']} (under half the median branch)"
                m["outliers"] = pd.DataFrame({"Bank": o["bank"], "Address": o["address"].astype(str),
                                              "Deposits": (o["deposits_usd"] / 1e6).map(lambda x: f"${x:,.1f}M"),
                                              "x median": o["vs_peer"].map(lambda x: f"{x:.2f}")})
    return m


def build_brief_docx(brief_md: str, label: str, sources: list[str], loaded_at: str) -> bytes:
    """Just the AI brief as a branded Word document (no scorecard tables)."""
    m = {"title": "Market Brief", "prepared_for": "", "purpose": label.split(" · ")[0], "date": datetime.now().strftime("%B %d, %Y").replace(" 0", " "),
         "data_as_of": loaded_at, "summary_md": brief_md, "settings": [], "scorecard": None, "sector_title": "", "sector_table": None,
         "banking": None, "competitors": None, "outliers": None, "sources": sources, "sample": False}
    return build_docx(m)


# ------------------------------------------------------------------ markdown-lite parsing
def _clean_inline(t: str) -> str:
    t = re.sub(r"`([^`]*)`", r"\1", t)
    t = re.sub(r"(?<!\*)\*(?!\*)([^*]+)(?<!\*)\*(?!\*)", r"\1", t)  # single-asterisk italics -> plain
    return t


def md_blocks(text: str) -> list[tuple]:
    out, lines, i = [], text.splitlines(), 0
    while i < len(lines):
        ln = lines[i].rstrip()
        s = ln.strip()
        if not s:
            i += 1
        elif s.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c):
                    rows.append([_clean_inline(c) for c in cells])
                i += 1
            if rows:
                out.append(("table", rows))
        elif re.match(r"#{1,4}\s", s):
            level = len(s) - len(s.lstrip("#"))
            out.append((f"h{min(level, 3)}", _clean_inline(s.lstrip("#").strip())))
            i += 1
        elif re.match(r"[-*•]\s+", s):
            items = []
            while i < len(lines) and re.match(r"\s*[-*•]\s+", lines[i]):
                items.append(_clean_inline(re.sub(r"^\s*[-*•]\s+", "", lines[i])))
                i += 1
            out.append(("ul", items))
        elif re.match(r"\d+[.)]\s+", s):
            items = []
            while i < len(lines) and re.match(r"\s*\d+[.)]\s+", lines[i]):
                items.append(_clean_inline(re.sub(r"^\s*\d+[.)]\s+", "", lines[i])))
                i += 1
            out.append(("ol", items))
        elif re.fullmatch(r"[-*_]{3,}", s):
            i += 1
        else:
            para = [s]
            i += 1
            while i < len(lines) and lines[i].strip() and not re.match(r"\s*([#|]|[-*•]\s|\d+[.)]\s)", lines[i]):
                para.append(lines[i].strip())
                i += 1
            out.append(("p", _clean_inline(" ".join(para))))
    return out


def _bold_parts(text: str) -> list[tuple[str, bool]]:
    parts = re.split(r"(\*\*[^*]+\*\*)", text)
    return [(p[2:-2], True) if p.startswith("**") and p.endswith("**") else (p, False) for p in parts if p]


# ------------------------------------------------------------------ Word
def build_docx(m: dict) -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor

    doc = Document()
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Inches(0.9)
    sec.top_margin, sec.bottom_margin = Inches(0.8), Inches(0.8)
    base = doc.styles["Normal"]
    base.font.name, base.font.size = "Calibri", Pt(10.5)

    def shade(cell, hex_):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hex_)
        tcPr.append(shd)

    def runs(par, text, size=None, color=None):
        for t, b in _bold_parts(text):
            r = par.add_run(t); r.bold = b
            if size: r.font.size = Pt(size)
            if color: r.font.color.rgb = RGBColor.from_string(color)

    def heading(text, level=1):
        p = doc.add_paragraph()
        p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(12 if level == 1 else 8), Pt(4)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(text); r.bold = True
        r.font.size = Pt({1: 14, 2: 12, 3: 11}[level]); r.font.color.rgb = RGBColor.from_string(NAVY if level == 1 else TEAL)
        if level == 1:  # thin rule under section headings
            pPr = p._p.get_or_add_pPr(); bd = OxmlElement("w:pBdr"); b = OxmlElement("w:bottom")
            b.set(qn("w:val"), "single"); b.set(qn("w:sz"), "6"); b.set(qn("w:space"), "1"); b.set(qn("w:color"), TEAL)
            bd.append(b); pPr.append(bd)

    def table(rows, header=True, widths=None):
        t = doc.add_table(rows=len(rows), cols=len(rows[0]))
        t.style = "Table Grid"
        for i, row in enumerate(rows):
            for j, val in enumerate(row):
                c = t.cell(i, j); c.text = ""
                p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(0)
                if i == 0 and header:
                    shade(c, NAVY); runs(p, str(val), 9, "FFFFFF"); [setattr(r, "bold", True) for r in p.runs]
                else:
                    if i % 2 == 0: shade(c, "EEF3F7")
                    runs(p, str(val), 9)
        doc.add_paragraph().paragraph_format.space_after = Pt(2)

    def df_table(df):
        table([list(df.columns)] + df.astype(str).values.tolist())

    # banner
    bt = doc.add_table(rows=1, cols=1); c = bt.cell(0, 0); shade(c, NAVY)
    p = c.paragraphs[0]; r = p.add_run("Synoptic BI"); r.bold = True; r.font.size = Pt(20); r.font.color.rgb = RGBColor(255, 255, 255)
    p2 = c.add_paragraph(); r2 = p2.add_run("FLORIDA MARKET INTELLIGENCE"); r2.font.size = Pt(8.5); r2.font.color.rgb = RGBColor(0xA5, 0xD8, 0xD4)
    doc.add_paragraph()
    t = doc.add_paragraph(); rt = t.add_run(m["title"]); rt.bold = True; rt.font.size = Pt(18); rt.font.color.rgb = RGBColor.from_string(NAVY)
    meta = f"{m['date']}  |  Purpose: {m['purpose']}  |  Data as of {m['data_as_of']}"
    if m["prepared_for"]:
        meta = f"Prepared for {m['prepared_for']}  |  " + meta
    mp = doc.add_paragraph(); rm = mp.add_run(meta); rm.font.size = Pt(9); rm.font.color.rgb = RGBColor(0x55, 0x65, 0x75)
    if m["sample"]:
        wp = doc.add_paragraph(); rw = wp.add_run("Note: some data sources were on illustrative sample values when this memo was generated.")
        rw.bold = True; rw.font.size = Pt(9); rw.font.color.rgb = RGBColor(0xB4, 0x53, 0x09)

    heading("1. Executive summary" if m.get("scorecard") is not None else "Brief")
    if m["summary_md"]:
        for kind, val in md_blocks(m["summary_md"]):
            if kind in ("h1", "h2", "h3"):
                heading(val, 2 if kind != "h3" else 3)
            elif kind == "p":
                p = doc.add_paragraph(); runs(p, val)
            elif kind in ("ul", "ol"):
                for it in val:
                    p = doc.add_paragraph(style="List Bullet" if kind == "ul" else "List Number"); runs(p, it)
            elif kind == "table":
                table(val)
    else:
        doc.add_paragraph("No AI-written summary was attached. Generate a brief in the AI Analyst tab, then export again "
                          "to include it here.")

    if m.get("scorecard") is not None:
        heading("2. Macro scorecard")
        doc.add_paragraph("Markets ranked by the Synoptic BI score, a relative 0-100 measure within this set of markets.")
        df_table(m["scorecard"])
        for s_ in m["settings"]:
            p = doc.add_paragraph(style="List Bullet"); runs(p, s_, 9.5)
        heading(m["sector_title"], 2)
        df_table(m["sector_table"])

    if m["banking"] is not None:
        heading("3. Micro spatial analysis: small businesses vs bank branches")
        df_table(m["banking"])
        p = doc.add_paragraph(); runs(p, m["banking_note"], 9)
        if m["competitors"] is not None:
            heading(m["competitors_title"], 2); df_table(m["competitors"])
        if m["outliers"] is not None:
            heading(m["outliers_title"], 2); df_table(m["outliers"])
    n = (3 if m.get("scorecard") is not None else 2) + (1 if m["banking"] is not None else 0)
    heading(f"{n}. Data sources and compliance" if m.get("scorecard") is not None else "Data sources and compliance")
    for s_ in m["sources"]:
        p = doc.add_paragraph(style="List Bullet"); runs(p, s_, 9)
    p = doc.add_paragraph(); runs(p, DISCLAIMER, 8.5)

    fp = sec.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = fp.add_run("Synoptic BI | Confidential screening memo | Page "); fr.font.size = Pt(8)
    for tag, txt in (("begin", None), (None, "PAGE"), ("end", None)):
        run = fp.add_run(); run.font.size = Pt(8)
        if tag:
            el = OxmlElement("w:fldChar"); el.set(qn("w:fldCharType"), tag); run._r.append(el)
        else:
            el = OxmlElement("w:instrText"); el.set(qn("xml:space"), "preserve"); el.text = txt; run._r.append(el)
    buf = io.BytesIO(); doc.save(buf)
    return buf.getvalue()


# ------------------------------------------------------------------ PDF
def _pdf_text(t: str) -> str:
    t = (str(t).replace("≥", ">=").replace("≤", "<=").replace("→", "->").replace("×", "x").replace("✅", "").replace("❌", "")
         .replace(" ", " "))
    return t.encode("cp1252", "ignore").decode("cp1252")


def build_pdf(m: dict) -> bytes:
    from xml.sax.saxutils import escape

    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import inch
    from reportlab.platypus import ListFlowable, ListItem, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    navy, teal = colors.HexColor("#" + NAVY), colors.HexColor("#" + TEAL)
    ss = getSampleStyleSheet()
    body = ParagraphStyle("body", parent=ss["BodyText"], fontName="Helvetica", fontSize=9.6, leading=13, spaceAfter=5)
    small = ParagraphStyle("small", parent=body, fontSize=8, leading=10.5, textColor=colors.HexColor("#4b5b6b"))
    h1 = ParagraphStyle("h1", parent=body, fontName="Helvetica-Bold", fontSize=13, textColor=navy, spaceBefore=12, spaceAfter=5)
    h2 = ParagraphStyle("h2", parent=body, fontName="Helvetica-Bold", fontSize=11, textColor=teal, spaceBefore=8, spaceAfter=4)
    title = ParagraphStyle("title", parent=body, fontName="Helvetica-Bold", fontSize=18, textColor=navy, spaceAfter=3, leading=22)
    cell = ParagraphStyle("cell", parent=body, fontSize=8, leading=9.8, spaceAfter=0)
    cell_h = ParagraphStyle("cellh", parent=cell, fontName="Helvetica-Bold", textColor=colors.white)

    def P(text, style=body):
        safe = escape(_pdf_text(text))
        safe = re.sub(r"\*\*([^*]+)\*\*", r"<b>\1</b>", safe)
        return Paragraph(safe, style)

    def tbl(rows):
        data = [[Paragraph(escape(_pdf_text(c)), cell_h if i == 0 else cell) for c in r] for i, r in enumerate(rows)]
        t = Table(data, repeatRows=1, hAlign="LEFT")
        t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), navy),
                               ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#EEF3F7")]),
                               ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#c4ced8")),
                               ("VALIGN", (0, 0), (-1, -1), "TOP"),
                               ("LEFTPADDING", (0, 0), (-1, -1), 4), ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                               ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3)]))
        return t

    def df_t(df):
        return tbl([list(df.columns)] + df.astype(str).values.tolist())

    def bullets(items, style=body):
        return ListFlowable([ListItem(P(i, style), leftIndent=12) for i in items], bulletType="bullet", start="•",
                            leftIndent=12)

    story = []
    meta = f"{m['date']}  |  Purpose: {m['purpose']}  |  Data as of {m['data_as_of']}"
    if m["prepared_for"]:
        meta = f"Prepared for {m['prepared_for']}  |  " + meta
    story += [Spacer(1, 0.55 * inch), P(m["title"], title), P(meta, small)]
    if m["sample"]:
        story.append(P("**Note:** some data sources were on illustrative sample values when this memo was generated.", small))

    story.append(P("1. Executive summary", h1))
    if m["summary_md"]:
        for kind, val in md_blocks(m["summary_md"]):
            if kind in ("h1", "h2", "h3"):
                story.append(P(val, h2))
            elif kind == "p":
                story.append(P(val))
            elif kind in ("ul", "ol"):
                story.append(bullets(val))
            elif kind == "table":
                story.append(tbl(val)); story.append(Spacer(1, 6))
    else:
        story.append(P("No AI-written summary was attached. Generate a brief in the AI Analyst tab, then export again."))

    story += [P("2. Macro scorecard", h1),
              P("Markets ranked by the Synoptic BI score, a relative 0-100 measure within this set of markets."),
              df_t(m["scorecard"]), Spacer(1, 4), bullets(m["settings"], small),
              P(m["sector_title"], h2), df_t(m["sector_table"])]
    if m["banking"] is not None:
        story += [P("3. Micro spatial analysis: small businesses vs bank branches", h1), df_t(m["banking"]),
                  Spacer(1, 3), P(m["banking_note"], small)]
        if m["competitors"] is not None:
            story += [P(m["competitors_title"], h2), df_t(m["competitors"])]
        if m["outliers"] is not None:
            story += [P(m["outliers_title"], h2), df_t(m["outliers"])]
    n = 4 if m["banking"] is not None else 3
    story += [P(f"{n}. Data sources and compliance", h1), bullets(m["sources"], small), Spacer(1, 4), P(DISCLAIMER, small)]

    def band(canvas, doc_):
        canvas.saveState()
        w, h = letter
        if doc_.page == 1:
            canvas.setFillColor(navy); canvas.rect(0, h - 0.95 * inch, w, 0.95 * inch, stroke=0, fill=1)
            canvas.setFillColor(colors.white); canvas.setFont("Helvetica-Bold", 20)
            canvas.drawString(0.8 * inch, h - 0.55 * inch, "Synoptic BI")
            canvas.setFillColor(colors.HexColor("#A5D8D4")); canvas.setFont("Helvetica", 8.5)
            canvas.drawString(0.8 * inch, h - 0.78 * inch, "FLORIDA MARKET INTELLIGENCE")
        canvas.setFillColor(colors.HexColor("#6b7a8a")); canvas.setFont("Helvetica", 8)
        canvas.drawCentredString(w / 2, 0.45 * inch, f"Synoptic BI | Confidential screening memo | Page {doc_.page}")
        canvas.restoreState()

    buf = io.BytesIO()
    SimpleDocTemplate(buf, pagesize=letter, leftMargin=0.8 * inch, rightMargin=0.8 * inch, topMargin=0.7 * inch,
                      bottomMargin=0.8 * inch, title=m["title"], author="Synoptic BI").build(
        story, onFirstPage=band, onLaterPages=band)
    return buf.getvalue()
